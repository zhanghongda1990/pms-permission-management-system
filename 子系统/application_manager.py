import json
import os
from datetime import datetime
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


class ApplicationManager:
    def __init__(self):
        self.package_applications = []
        self.authorization_applications = []
        
    def create_package_application(self, name, description, permissions, document_no='', applicant='', department='', organization='', operation_type='new', original_permissions=None):
        application = {
            'id': self._generate_id(),
            'type': 'package_application',
            'name': name,
            'description': description,
            'permissions': permissions,
            'document_no': document_no,
            'applicant': applicant,
            'department': department,
            'organization': organization,
            'operation_type': operation_type,
            'original_permissions': original_permissions or [],
            'create_time': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
            'status': 'pending'
        }
        self.package_applications.append(application)
        return application
    
    def create_authorization_application(self, personnel_info, packages, document_no='', applicant='', department='', organization=''):
        all_permissions = []
        for pkg in packages:
            for perm in pkg.get('permissions', []):
                all_permissions.append(perm)
                
        application = {
            'id': self._generate_id(),
            'type': 'authorization_application',
            'personnel': personnel_info,
            'packages': packages,
            'permissions': all_permissions,
            'document_no': document_no,
            'applicant': applicant,
            'department': department,
            'organization': organization,
            'create_time': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
            'status': 'pending'
        }
        self.authorization_applications.append(application)
        return application
    
    def create_authorization_change_application(self, personnel_info, packages_to_add, packages_to_remove, document_no='', applicant='', department='', organization=''):
        permissions_to_add = []
        for pkg in packages_to_add:
            for perm in pkg.get('permissions', []):
                permissions_to_add.append(perm)
                
        permissions_to_remove = []
        for pkg in packages_to_remove:
            for perm in pkg.get('permissions', []):
                permissions_to_remove.append(perm)
                
        application = {
            'id': self._generate_id(),
            'type': 'authorization_change_application',
            'personnel': personnel_info,
            'packages_to_add': packages_to_add,
            'packages_to_remove': packages_to_remove,
            'permissions_to_add': permissions_to_add,
            'permissions_to_remove': permissions_to_remove,
            'document_no': document_no,
            'applicant': applicant,
            'department': department,
            'organization': organization,
            'create_time': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
            'status': 'pending'
        }
        self.authorization_applications.append(application)
        return application
    
    def _generate_id(self):
        import uuid
        return str(uuid.uuid4())[:8].upper()
    
    def export_package_application_yml(self, application, output_path):
        operation_type = application.get('operation_type', 'new')
        original_permissions = application.get('original_permissions', [])
        current_permissions = application.get('permissions', [])
        
        added_perms, removed_perms = self._compare_permissions(original_permissions, current_permissions)
        
        export_data = {
            'export_type': '权限包申请',
            'export_time': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
            'application': {
                'id': application['id'],
                'name': application['name'],
                'description': application.get('description', ''),
                'document_no': application.get('document_no', ''),
                'applicant': application.get('applicant', ''),
                'department': application.get('department', ''),
                'organization': application.get('organization', ''),
                'operation_type': operation_type,
                'create_time': application['create_time'],
                'status': application['status']
            }
        }
        
        if operation_type == 'new':
            export_data['application']['permissions'] = current_permissions
            export_data['application']['permission_count'] = len(current_permissions)
        else:
            export_data['application']['original_permissions'] = original_permissions
            export_data['application']['current_permissions'] = current_permissions
            export_data['application']['added_permissions'] = added_perms
            export_data['application']['removed_permissions'] = removed_perms
            export_data['application']['added_count'] = len(added_perms)
            export_data['application']['removed_count'] = len(removed_perms)
        
        with open(output_path, 'w', encoding='utf-8') as f:
            json.dump(export_data, f, ensure_ascii=False, indent=2)
        
        return output_path
    
    def export_authorization_application_yml(self, application, output_path):
        export_data = {
            'export_type': '授权申请',
            'export_time': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
            'application': {
                'id': application['id'],
                'personnel': application['personnel'],
                'packages': [{'name': pkg['name'], 'permissions': pkg.get('permissions', [])} for pkg in application.get('packages', [])],
                'permissions': application['permissions'],
                'document_no': application.get('document_no', ''),
                'applicant': application.get('applicant', ''),
                'department': application.get('department', ''),
                'organization': application.get('organization', ''),
                'create_time': application['create_time'],
                'status': application['status']
            }
        }
        
        with open(output_path, 'w', encoding='utf-8') as f:
            json.dump(export_data, f, ensure_ascii=False, indent=2)
        
        return output_path
    
    def export_authorization_change_application_yml(self, application, output_path):
        export_data = {
            'export_type': '授权变更申请',
            'export_time': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
            'application': {
                'id': application['id'],
                'type': 'authorization_change',
                'personnel': application['personnel'],
                'packages_to_add': [{'name': pkg['name'], 'permissions': pkg.get('permissions', [])} for pkg in application.get('packages_to_add', [])],
                'packages_to_remove': [{'name': pkg['name'], 'permissions': pkg.get('permissions', [])} for pkg in application.get('packages_to_remove', [])],
                'permissions_to_add': application.get('permissions_to_add', []),
                'permissions_to_remove': application.get('permissions_to_remove', []),
                'document_no': application.get('document_no', ''),
                'applicant': application.get('applicant', ''),
                'department': application.get('department', ''),
                'organization': application.get('organization', ''),
                'create_time': application['create_time'],
                'status': application['status']
            }
        }
        
        with open(output_path, 'w', encoding='utf-8') as f:
            json.dump(export_data, f, ensure_ascii=False, indent=2)
        
        return output_path
    
    def export_package_application_doc(self, application, output_path):
        doc = Document()
        
        self._set_document_style(doc)
        
        org_para = doc.add_paragraph()
        org_run = org_para.add_run(f"申请机构：{application.get('organization', '')}")
        org_run.font.size = Pt(14)
        org_run.font.bold = True
        org_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_paragraph()
        
        title = doc.add_paragraph()
        title_run = title.add_run('权限包申请表')
        title_run.font.size = Pt(18)
        title_run.font.bold = True
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_paragraph()
        
        operation_type = application.get('operation_type', 'new')
        operation_text = '新建权限包' if operation_type == 'new' else '变更权限包'
        
        info_table = doc.add_table(rows=5, cols=4)
        info_table.style = 'Table Grid'
        
        self._set_cell_text(info_table.cell(0, 0), '申请编号')
        self._set_cell_text(info_table.cell(0, 1), application['id'])
        self._set_cell_text(info_table.cell(0, 2), '申请日期')
        self._set_cell_text(info_table.cell(0, 3), application['create_time'][:10])
        
        self._set_cell_text(info_table.cell(1, 0), '权限包名称')
        self._set_cell_text(info_table.cell(1, 1), application['name'], colspan=3)
        info_table.cell(1, 1).merge(info_table.cell(1, 3))
        
        self._set_cell_text(info_table.cell(2, 0), '操作类型')
        self._set_cell_text(info_table.cell(2, 1), operation_text, colspan=3)
        info_table.cell(2, 1).merge(info_table.cell(2, 3))
        
        self._set_cell_text(info_table.cell(3, 0), '申请人')
        self._set_cell_text(info_table.cell(3, 1), application.get('applicant', ''))
        self._set_cell_text(info_table.cell(3, 2), '所属部门')
        self._set_cell_text(info_table.cell(3, 3), application.get('department', ''))
        
        self._set_cell_text(info_table.cell(4, 0), '文档编号')
        self._set_cell_text(info_table.cell(4, 1), application.get('document_no', ''))
        self._set_cell_text(info_table.cell(4, 2), '状态')
        self._set_cell_text(info_table.cell(4, 3), '待审批')
        
        doc.add_paragraph()
        
        if application.get('description'):
            desc_title = doc.add_paragraph()
            desc_title.add_run('权限包描述：').bold = True
            doc.add_paragraph(application['description'])
            doc.add_paragraph()
        
        permissions = application['permissions']
        original_permissions = application.get('original_permissions', [])
        
        if operation_type == 'new':
            perm_title = doc.add_paragraph()
            perm_title.add_run('新建权限包包含权限：').bold = True
            
            if permissions:
                perm_table = doc.add_table(rows=1, cols=5)
                perm_table.style = 'Table Grid'
                
                headers = ['序号', '系统名称', '权限名称', '权限代码', '权限级别']
                for i, header in enumerate(headers):
                    self._set_cell_text(perm_table.cell(0, i), header, bold=True)
                
                for idx, perm in enumerate(permissions, 1):
                    row = perm_table.add_row()
                    self._set_cell_text(row.cells[0], str(idx))
                    self._set_cell_text(row.cells[1], perm.get('system', ''))
                    self._set_cell_text(row.cells[2], perm.get('permission_name', ''))
                    self._set_cell_text(row.cells[3], perm.get('permission_code', ''))
                    self._set_cell_text(row.cells[4], str(perm.get('permission_level', '')))
        else:
            added_perms, removed_perms = self._compare_permissions(original_permissions, permissions)
            
            if added_perms:
                add_title = doc.add_paragraph()
                add_title.add_run('新增权限：').bold = True
                
                add_table = doc.add_table(rows=1, cols=5)
                add_table.style = 'Table Grid'
                
                headers = ['序号', '系统名称', '权限名称', '权限代码', '权限级别']
                for i, header in enumerate(headers):
                    self._set_cell_text(add_table.cell(0, i), header, bold=True)
                
                for idx, perm in enumerate(added_perms, 1):
                    row = add_table.add_row()
                    self._set_cell_text(row.cells[0], str(idx))
                    self._set_cell_text(row.cells[1], perm.get('system', ''))
                    self._set_cell_text(row.cells[2], perm.get('permission_name', ''))
                    self._set_cell_text(row.cells[3], perm.get('permission_code', ''))
                    self._set_cell_text(row.cells[4], str(perm.get('permission_level', '')))
                
                doc.add_paragraph()
            
            if removed_perms:
                remove_title = doc.add_paragraph()
                remove_title.add_run('移除权限：').bold = True
                
                remove_table = doc.add_table(rows=1, cols=5)
                remove_table.style = 'Table Grid'
                
                headers = ['序号', '系统名称', '权限名称', '权限代码', '权限级别']
                for i, header in enumerate(headers):
                    self._set_cell_text(remove_table.cell(0, i), header, bold=True)
                
                for idx, perm in enumerate(removed_perms, 1):
                    row = remove_table.add_row()
                    self._set_cell_text(row.cells[0], str(idx))
                    self._set_cell_text(row.cells[1], perm.get('system', ''))
                    self._set_cell_text(row.cells[2], perm.get('permission_name', ''))
                    self._set_cell_text(row.cells[3], perm.get('permission_code', ''))
                    self._set_cell_text(row.cells[4], str(perm.get('permission_level', '')))
        
        doc.add_paragraph()
        doc.add_paragraph()
        
        sign_title = doc.add_paragraph()
        sign_title.add_run('签字确认：').bold = True
        
        sign_table = doc.add_table(rows=2, cols=2)
        sign_table.style = 'Table Grid'
        
        self._set_cell_text(sign_table.cell(0, 0), '申请人签字')
        self._set_cell_text(sign_table.cell(0, 1), '审批人签字')
        
        self._set_cell_text(sign_table.cell(1, 0), '\n\n\n日期：')
        self._set_cell_text(sign_table.cell(1, 1), '\n\n\n日期：')
        
        self._add_page_numbers(doc)
        
        doc.save(output_path)
        return output_path
    
    def _compare_permissions(self, original, current):
        original_keys = {f"{p.get('system', '')}_{p.get('permission_code', '')}" for p in original}
        current_keys = {f"{p.get('system', '')}_{p.get('permission_code', '')}" for p in current}
        
        added_keys = current_keys - original_keys
        removed_keys = original_keys - current_keys
        
        added_perms = [p for p in current if f"{p.get('system', '')}_{p.get('permission_code', '')}" in added_keys]
        removed_perms = [p for p in original if f"{p.get('system', '')}_{p.get('permission_code', '')}" in removed_keys]
        
        return added_perms, removed_perms
    
    def export_authorization_application_doc(self, application, output_path):
        doc = Document()
        
        self._set_document_style(doc)
        
        org_para = doc.add_paragraph()
        org_run = org_para.add_run(f"申请机构：{application.get('organization', '')}")
        org_run.font.size = Pt(14)
        org_run.font.bold = True
        org_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_paragraph()
        
        title = doc.add_paragraph()
        title_run = title.add_run('权限授权申请表')
        title_run.font.size = Pt(18)
        title_run.font.bold = True
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_paragraph()
        
        personnel = application['personnel']
        
        info_table = doc.add_table(rows=7, cols=4)
        info_table.style = 'Table Grid'
        
        self._set_cell_text(info_table.cell(0, 0), '申请编号')
        self._set_cell_text(info_table.cell(0, 1), application['id'])
        self._set_cell_text(info_table.cell(0, 2), '申请日期')
        self._set_cell_text(info_table.cell(0, 3), application['create_time'][:10])
        
        self._set_cell_text(info_table.cell(1, 0), '被授权人')
        self._set_cell_text(info_table.cell(1, 1), personnel.get('name', ''))
        self._set_cell_text(info_table.cell(1, 2), '工号')
        self._set_cell_text(info_table.cell(1, 3), personnel.get('employee_id', ''))
        
        self._set_cell_text(info_table.cell(2, 0), '证件号码')
        self._set_cell_text(info_table.cell(2, 1), personnel.get('id_card', ''))
        self._set_cell_text(info_table.cell(2, 2), '电话号码')
        self._set_cell_text(info_table.cell(2, 3), personnel.get('phone', ''))
        
        self._set_cell_text(info_table.cell(3, 0), '所属部门')
        self._set_cell_text(info_table.cell(3, 1), personnel.get('department', ''))
        self._set_cell_text(info_table.cell(3, 2), '职位')
        self._set_cell_text(info_table.cell(3, 3), personnel.get('position', ''))
        
        self._set_cell_text(info_table.cell(4, 0), '申请人')
        self._set_cell_text(info_table.cell(4, 1), application.get('applicant', ''))
        self._set_cell_text(info_table.cell(4, 2), '申请人部门')
        self._set_cell_text(info_table.cell(4, 3), application.get('department', ''))
        
        self._set_cell_text(info_table.cell(5, 0), '文档编号')
        self._set_cell_text(info_table.cell(5, 1), application.get('document_no', ''))
        self._set_cell_text(info_table.cell(5, 2), '状态')
        self._set_cell_text(info_table.cell(5, 3), '待审批')
        
        self._set_cell_text(info_table.cell(6, 0), '备注')
        self._set_cell_text(info_table.cell(6, 1), personnel.get('remark', ''), colspan=3)
        info_table.cell(6, 1).merge(info_table.cell(6, 3))
        
        doc.add_paragraph()
        
        packages = application.get('packages', [])
        if packages:
            pkg_title = doc.add_paragraph()
            pkg_title.add_run('授权权限包：').bold = True
            
            pkg_table = doc.add_table(rows=1, cols=3)
            pkg_table.style = 'Table Grid'
            
            headers = ['序号', '权限包名称', '权限数量']
            for i, header in enumerate(headers):
                self._set_cell_text(pkg_table.cell(0, i), header, bold=True)
            
            for idx, pkg in enumerate(packages, 1):
                row = pkg_table.add_row()
                self._set_cell_text(row.cells[0], str(idx))
                self._set_cell_text(row.cells[1], pkg.get('name', ''))
                self._set_cell_text(row.cells[2], str(len(pkg.get('permissions', []))))
            
            doc.add_paragraph()
        
        perm_title = doc.add_paragraph()
        perm_title.add_run('授权权限明细：').bold = True
        
        permissions = application['permissions']
        if permissions:
            perm_table = doc.add_table(rows=1, cols=5)
            perm_table.style = 'Table Grid'
            
            headers = ['序号', '系统名称', '权限名称', '权限代码', '权限级别']
            for i, header in enumerate(headers):
                self._set_cell_text(perm_table.cell(0, i), header, bold=True)
            
            for idx, perm in enumerate(permissions, 1):
                row = perm_table.add_row()
                self._set_cell_text(row.cells[0], str(idx))
                self._set_cell_text(row.cells[1], perm.get('system', ''))
                self._set_cell_text(row.cells[2], perm.get('permission_name', ''))
                self._set_cell_text(row.cells[3], perm.get('permission_code', ''))
                self._set_cell_text(row.cells[4], str(perm.get('permission_level', '')))
        
        doc.add_paragraph()
        doc.add_paragraph()
        
        sign_title = doc.add_paragraph()
        sign_title.add_run('签字确认：').bold = True
        
        sign_table = doc.add_table(rows=2, cols=2)
        sign_table.style = 'Table Grid'
        
        self._set_cell_text(sign_table.cell(0, 0), '申请人签字')
        self._set_cell_text(sign_table.cell(0, 1), '审批人签字')
        
        self._set_cell_text(sign_table.cell(1, 0), '\n\n\n日期：')
        self._set_cell_text(sign_table.cell(1, 1), '\n\n\n日期：')
        
        self._add_page_numbers(doc)
        
        doc.save(output_path)
        return output_path
    
    def export_authorization_change_application_doc(self, application, output_path):
        doc = Document()
        
        self._set_document_style(doc)
        
        org_para = doc.add_paragraph()
        org_run = org_para.add_run(f"申请机构：{application.get('organization', '')}")
        org_run.font.size = Pt(14)
        org_run.font.bold = True
        org_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_paragraph()
        
        title = doc.add_paragraph()
        title_run = title.add_run('权限授权变更申请表')
        title_run.font.size = Pt(18)
        title_run.font.bold = True
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_paragraph()
        
        personnel = application['personnel']
        
        info_table = doc.add_table(rows=7, cols=4)
        info_table.style = 'Table Grid'
        
        self._set_cell_text(info_table.cell(0, 0), '申请编号')
        self._set_cell_text(info_table.cell(0, 1), application['id'])
        self._set_cell_text(info_table.cell(0, 2), '申请日期')
        self._set_cell_text(info_table.cell(0, 3), application['create_time'][:10])
        
        self._set_cell_text(info_table.cell(1, 0), '被授权人')
        self._set_cell_text(info_table.cell(1, 1), personnel.get('name', ''))
        self._set_cell_text(info_table.cell(1, 2), '工号')
        self._set_cell_text(info_table.cell(1, 3), personnel.get('employee_id', ''))
        
        self._set_cell_text(info_table.cell(2, 0), '证件号码')
        self._set_cell_text(info_table.cell(2, 1), personnel.get('id_card', ''))
        self._set_cell_text(info_table.cell(2, 2), '电话号码')
        self._set_cell_text(info_table.cell(2, 3), personnel.get('phone', ''))
        
        self._set_cell_text(info_table.cell(3, 0), '所属部门')
        self._set_cell_text(info_table.cell(3, 1), personnel.get('department', ''))
        self._set_cell_text(info_table.cell(3, 2), '职位')
        self._set_cell_text(info_table.cell(3, 3), personnel.get('position', ''))
        
        self._set_cell_text(info_table.cell(4, 0), '申请人')
        self._set_cell_text(info_table.cell(4, 1), application.get('applicant', ''))
        self._set_cell_text(info_table.cell(4, 2), '申请人部门')
        self._set_cell_text(info_table.cell(4, 3), application.get('department', ''))
        
        self._set_cell_text(info_table.cell(5, 0), '文档编号')
        self._set_cell_text(info_table.cell(5, 1), application.get('document_no', ''))
        self._set_cell_text(info_table.cell(5, 2), '状态')
        self._set_cell_text(info_table.cell(5, 3), '待审批')
        
        self._set_cell_text(info_table.cell(6, 0), '备注')
        self._set_cell_text(info_table.cell(6, 1), personnel.get('remark', ''), colspan=3)
        info_table.cell(6, 1).merge(info_table.cell(6, 3))
        
        doc.add_paragraph()
        
        packages_to_add = application.get('packages_to_add', [])
        if packages_to_add:
            add_title = doc.add_paragraph()
            add_title.add_run('新增授权权限包：').bold = True
            
            pkg_table = doc.add_table(rows=1, cols=3)
            pkg_table.style = 'Table Grid'
            
            headers = ['序号', '权限包名称', '权限数量']
            for i, header in enumerate(headers):
                self._set_cell_text(pkg_table.cell(0, i), header, bold=True)
            
            for idx, pkg in enumerate(packages_to_add, 1):
                row = pkg_table.add_row()
                self._set_cell_text(row.cells[0], str(idx))
                self._set_cell_text(row.cells[1], pkg.get('name', ''))
                self._set_cell_text(row.cells[2], str(len(pkg.get('permissions', []))))
            
            doc.add_paragraph()
        
        packages_to_remove = application.get('packages_to_remove', [])
        if packages_to_remove:
            remove_title = doc.add_paragraph()
            remove_title.add_run('移除授权权限包：').bold = True
            
            pkg_table = doc.add_table(rows=1, cols=3)
            pkg_table.style = 'Table Grid'
            
            headers = ['序号', '权限包名称', '权限数量']
            for i, header in enumerate(headers):
                self._set_cell_text(pkg_table.cell(0, i), header, bold=True)
            
            for idx, pkg in enumerate(packages_to_remove, 1):
                row = pkg_table.add_row()
                self._set_cell_text(row.cells[0], str(idx))
                self._set_cell_text(row.cells[1], pkg.get('name', ''))
                self._set_cell_text(row.cells[2], str(len(pkg.get('permissions', []))))
            
            doc.add_paragraph()
        
        permissions_to_add = application.get('permissions_to_add', [])
        if permissions_to_add:
            perm_title = doc.add_paragraph()
            perm_title.add_run('新增授权权限明细：').bold = True
            
            perm_table = doc.add_table(rows=1, cols=5)
            perm_table.style = 'Table Grid'
            
            headers = ['序号', '系统名称', '权限名称', '权限代码', '权限级别']
            for i, header in enumerate(headers):
                self._set_cell_text(perm_table.cell(0, i), header, bold=True)
            
            for idx, perm in enumerate(permissions_to_add, 1):
                row = perm_table.add_row()
                self._set_cell_text(row.cells[0], str(idx))
                self._set_cell_text(row.cells[1], perm.get('system', ''))
                self._set_cell_text(row.cells[2], perm.get('permission_name', ''))
                self._set_cell_text(row.cells[3], perm.get('permission_code', ''))
                self._set_cell_text(row.cells[4], str(perm.get('permission_level', '')))
            
            doc.add_paragraph()
        
        permissions_to_remove = application.get('permissions_to_remove', [])
        if permissions_to_remove:
            perm_title = doc.add_paragraph()
            perm_title.add_run('移除授权权限明细：').bold = True
            
            perm_table = doc.add_table(rows=1, cols=5)
            perm_table.style = 'Table Grid'
            
            headers = ['序号', '系统名称', '权限名称', '权限代码', '权限级别']
            for i, header in enumerate(headers):
                self._set_cell_text(perm_table.cell(0, i), header, bold=True)
            
            for idx, perm in enumerate(permissions_to_remove, 1):
                row = perm_table.add_row()
                self._set_cell_text(row.cells[0], str(idx))
                self._set_cell_text(row.cells[1], perm.get('system', ''))
                self._set_cell_text(row.cells[2], perm.get('permission_name', ''))
                self._set_cell_text(row.cells[3], perm.get('permission_code', ''))
                self._set_cell_text(row.cells[4], str(perm.get('permission_level', '')))
        
        doc.add_paragraph()
        doc.add_paragraph()
        
        sign_title = doc.add_paragraph()
        sign_title.add_run('签字确认：').bold = True
        
        sign_table = doc.add_table(rows=2, cols=2)
        sign_table.style = 'Table Grid'
        
        self._set_cell_text(sign_table.cell(0, 0), '申请人签字')
        self._set_cell_text(sign_table.cell(0, 1), '审批人签字')
        
        self._set_cell_text(sign_table.cell(1, 0), '\n\n\n日期：')
        self._set_cell_text(sign_table.cell(1, 1), '\n\n\n日期：')
        
        self._add_page_numbers(doc)
        
        doc.save(output_path)
        return output_path
    
    def _set_document_style(self, doc):
        doc.styles['Normal'].font.name = '宋体'
        doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        doc.styles['Normal'].font.size = Pt(11)
    
    def _set_cell_text(self, cell, text, bold=False, colspan=1):
        cell.text = text
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.name = '宋体'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
                run.font.size = Pt(10)
                if bold:
                    run.font.bold = True
    
    def _add_page_numbers(self, doc):
        sectPr = doc.sections[0]._sectPr
        pgNumType = OxmlElement('w:pgNumType')
        pgNumType.set(qn('w:start'), '1')
        sectPr.append(pgNumType)
        
        footer = doc.sections[0].footer
        footer.is_linked_to_previous = False
        
        footer_para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        run1 = footer_para.add_run('第 ')
        run1.font.name = '宋体'
        run1._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        run1.font.size = Pt(10)
        
        fldChar1 = OxmlElement('w:fldChar')
        fldChar1.set(qn('w:fldCharType'), 'begin')
        
        instrText = OxmlElement('w:instrText')
        instrText.text = "PAGE"
        
        fldChar2 = OxmlElement('w:fldChar')
        fldChar2.set(qn('w:fldCharType'), 'end')
        
        run2 = footer_para.add_run()
        run2._r.append(fldChar1)
        run2._r.append(instrText)
        run2._r.append(fldChar2)
        run2.font.name = '宋体'
        run2._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        run2.font.size = Pt(10)
        
        run3 = footer_para.add_run(' 页 共 ')
        run3.font.name = '宋体'
        run3._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        run3.font.size = Pt(10)
        
        fldChar3 = OxmlElement('w:fldChar')
        fldChar3.set(qn('w:fldCharType'), 'begin')
        
        instrText2 = OxmlElement('w:instrText')
        instrText2.text = "NUMPAGES"
        
        fldChar4 = OxmlElement('w:fldChar')
        fldChar4.set(qn('w:fldCharType'), 'end')
        
        run4 = footer_para.add_run()
        run4._r.append(fldChar3)
        run4._r.append(instrText2)
        run4._r.append(fldChar4)
        run4.font.name = '宋体'
        run4._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        run4.font.size = Pt(10)
        
        run5 = footer_para.add_run(' 页')
        run5.font.name = '宋体'
        run5._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        run5.font.size = Pt(10)
